#!/usr/bin/env python3
"""Generate PPT Presentasi UTS - RESIK Bank Sampah Digital"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_R = os.path.join(BASE_DIR, "..", "R.png")
LOGO_RESIK = os.path.join(BASE_DIR, "..", "Resik.png")

GREEN_DARK = RGBColor(0x1B, 0x5E, 0x20)
GREEN_MAIN = RGBColor(0x2E, 0x7D, 0x32)
GREEN_LIGHT = RGBColor(0x4C, 0xAF, 0x50)
GREEN_BG = RGBColor(0xE8, 0xF5, 0xE9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x1A)
GRAY_TEXT = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x00, 0xC8, 0x53)
ORANGE = RGBColor(0xFF, 0x8F, 0x00)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height

def add_bg(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = color
    bg.line.fill.background(); bg.shadow.inherit = False
    sp = bg._element; sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)

def add_shape(slide, st, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(st, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line: s.line.color.rgb = line; s.line.width = Pt(1)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s

def add_text(slide, x, y, w, h, text, sz=18, bold=False, color=DARK_TEXT, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(sz); p.font.bold = bold; p.font.color.rgb = color
    p.font.name = "Segoe UI"; p.alignment = align
    return tb

def add_multi(slide, x, y, w, h, lines, sz=16, color=DARK_TEXT, align=PP_ALIGN.LEFT, sp=8):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, ld in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(ld, tuple): t, b = ld
        else: t, b = ld, False
        p.text = t; p.font.size = Pt(sz); p.font.bold = b
        p.font.color.rgb = color; p.font.name = "Segoe UI"
        p.alignment = align; p.space_after = Pt(sp)
    return tb

def logo_r(slide, x, y, w=Inches(0.5), h=Inches(0.5)):
    if os.path.exists(LOGO_R): slide.shapes.add_picture(LOGO_R, x, y, w, h)

def logo_resik(slide, x, y, w=Inches(1.5), h=Inches(0.5)):
    if os.path.exists(LOGO_RESIK): slide.shapes.add_picture(LOGO_RESIK, x, y, w, h)

def header(slide, title):
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.0), GREEN_DARK)
    logo_r(slide, Inches(0.3), Inches(0.22), Inches(0.55), Inches(0.55))
    add_text(slide, Inches(1.0), Inches(0.15), Inches(10), Inches(0.7), title, sz=24, bold=True, color=WHITE)
    logo_resik(slide, Inches(11.5), Inches(0.25), Inches(1.5), Inches(0.5))

def footer(slide, num):
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(7.2), SLIDE_W, Inches(0.3), GREEN_DARK)
    add_text(slide, Inches(0.3), Inches(7.22), Inches(6), Inches(0.25), "RESIK - Bank Sampah Digital | PBO UTS", sz=9, color=WHITE)
    add_text(slide, Inches(12), Inches(7.22), Inches(1), Inches(0.25), str(num), sz=9, color=WHITE, align=PP_ALIGN.RIGHT)

def card(slide, x, y, w, h, fill=GREEN_BG):
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)

def icon_circle(slide, x, y, d, text, bg=GREEN_MAIN, tc=WHITE, sz=20):
    c = add_shape(slide, MSO_SHAPE.OVAL, x, y, d, d, bg)
    tf = c.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(sz); p.font.bold = True; p.font.color.rgb = tc
    p.alignment = PP_ALIGN.CENTER
    return c

# ═══ SLIDE 1: COVER ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, GREEN_DARK)
add_shape(s, MSO_SHAPE.OVAL, Inches(-2), Inches(-2), Inches(6), Inches(6), GREEN_MAIN)
add_shape(s, MSO_SHAPE.OVAL, Inches(10), Inches(4), Inches(7), Inches(7), GREEN_MAIN)
add_shape(s, MSO_SHAPE.OVAL, Inches(9.5), Inches(0.3), Inches(4), Inches(4), GREEN_LIGHT)
logo_resik(s, Inches(5.4), Inches(0.5), Inches(2.5), Inches(0.9))
logo_r(s, Inches(5.9), Inches(1.8), Inches(1.5), Inches(1.5))
add_text(s, Inches(1), Inches(3.6), Inches(11.333), Inches(1), "RESIK", sz=52, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(4.4), Inches(11.333), Inches(0.8), "Bank Sampah Digital", sz=28, color=GREEN_LIGHT, align=PP_ALIGN.CENTER)
add_shape(s, MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(5.2), Inches(2.3), Inches(0.04), ACCENT)
add_text(s, Inches(1), Inches(5.4), Inches(11.333), Inches(0.5), "Presentasi Ujian Tengah Semester (UTS)", sz=18, color=WHITE, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(5.9), Inches(11.333), Inches(0.5), "Pemrograman Berorientasi Objek (PBO)", sz=16, color=GRAY_TEXT, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(6.4), Inches(11.333), Inches(0.5), "Kelompok: [Anggota 1] - [NIM]  |  [Anggota 2] - [NIM]  |  [Anggota 3] - [NIM]", sz=12, color=GRAY_TEXT, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(6.9), Inches(11.333), Inches(0.4), "Apache NetBeans | Java OOP | 2026", sz=11, color=GRAY_TEXT, align=PP_ALIGN.CENTER)

# ═══ SLIDE 2: AGENDA ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Agenda Presentasi")
agenda = [("1","Judul & Nama Proyek"),("2","Anggota Kelompok"),("3","Permasalahan Nyata"),("4","Alasan Proyek Dikembangkan"),("5","Tujuan Pengembangan"),("6","Target Pengguna & Manfaat"),("7","Fitur Utama"),("8","Alur Penggunaan Aplikasi"),("9","Rencana OOP"),("10","Tahapan Pengembangan hingga UAS"),("11","Pembagian Tugas Anggota"),("12","Target & Fitur Tambahan UAS")]
c1x, c2x = Inches(1.0), Inches(7.0)
sy, rh = Inches(1.3), Inches(0.48)
for i, (num, title) in enumerate(agenda):
    col, row = i // 6, i % 6
    x = c1x if col == 0 else c2x
    y = sy + row * rh
    icon_circle(s, x, y, Inches(0.4), num, sz=12)
    add_text(s, x + Inches(0.55), y + Inches(0.03), Inches(5), Inches(0.4), title, sz=14)
footer(s, 2)

# ═══ SLIDE 3: JUDUL PROYEK ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Judul & Nama Proyek")
card(s, Inches(0.8), Inches(1.5), Inches(5), Inches(5.2))
logo_r(s, Inches(2.3), Inches(2.0), Inches(2), Inches(2))
logo_resik(s, Inches(1.8), Inches(4.3), Inches(3), Inches(1))
add_text(s, Inches(1.0), Inches(5.5), Inches(4.6), Inches(0.8), "RESIK - Bank Sampah Digital", sz=18, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
add_multi(s, Inches(6.2), Inches(1.5), Inches(6.5), Inches(5), [("Nama Proyek:", True),("RESIK (Bank Sampah Digital)", False),("", False),("Platform:", True),("- Mobile App: masyarakat (penyetor sampah)", False),("- Desktop App: petugas (verifikator)", False),("", False),("Bahasa Pemrograman:", True),("Java (Apache NetBeans)", False),("", False),("Paradigma:", True),("Object-Oriented Programming (OOP)", False),("", False),("Tahap:", True),("UTS - Prototype & Rencana  |  UAS - Aplikasi Lengkap", False)], sz=14, sp=4)
footer(s, 3)

# ═══ SLIDE 4: ANGGOTA KELOMPOK ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Anggota Kelompok")
members = [("1","[Nama Mahasiswa]","[NIM]","Arsitektur & Backend Core"),("2","[Nama Mahasiswa]","[NIM]","Fitur Transaksi & Dompet Digital"),("3","[Nama Mahasiswa]","[NIM]","UI/UX & Edukasi Sampah")]
cw, ch, gap = Inches(3.7), Inches(4.8), Inches(0.3)
tw = cw * 3 + gap * 2; sx = (SLIDE_W - tw) // 2; cy = Inches(1.4)
for i, (num, name, nim, task) in enumerate(members):
    x = sx + i * (cw + gap)
    card(s, x, cy, cw, ch)
    icon_circle(s, x + (cw - Inches(1.2)) // 2, cy + Inches(0.3), Inches(1.2), num, sz=28)
    add_text(s, x + Inches(0.2), cy + Inches(1.7), cw - Inches(0.4), Inches(0.5), name, sz=16, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), cy + Inches(2.2), cw - Inches(0.4), Inches(0.4), nim, sz=14, color=GRAY_TEXT, align=PP_ALIGN.CENTER)
    add_shape(s, MSO_SHAPE.RECTANGLE, x + Inches(0.5), cy + Inches(2.8), cw - Inches(1.0), Inches(0.03), ACCENT)
    add_text(s, x + Inches(0.2), cy + Inches(3.0), cw - Inches(0.4), Inches(0.4), "Tugas:", sz=12, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.3), cy + Inches(3.4), cw - Inches(0.6), Inches(1.2), task, sz=13, align=PP_ALIGN.CENTER)
footer(s, 4)

# ═══ SLIDE 5: PERMASALAHAN ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Permasalahan Nyata")
problems = [("1","Minimnya Edukasi Pemilahan Sampah","Masyarakat tidak mengetahui jenis sampah yang dapat didaur ulang, dijual, atau dibuang. Pencampuran sampah organik & anorganik masih terjadi."),("2","Belum Ada Nilai Ekonomis dari Sampah","Banyak masyarakat tidak menyadari bahwa sampah yang dipilah dengan baik memiliki nilai jual dan dapat menjadi penghasilan tambahan."),("3","Proses Penyetoran Tidak Terstruktur","Masyarakat bingung ke mana harus menyetor, bagaimana prosesnya, dan berapa nilai yang akan didapat."),("4","Kurangnya Transparansi & Tracking","Tidak ada sistem yang memungkinkan masyarakat melacak riwayat setoran sampah beserta pendapatan yang diperoleh.")]
cw, ch = Inches(5.8), Inches(2.5); gx, gy = Inches(0.3), Inches(0.3)
sx2, sy2 = Inches(0.8), Inches(1.3)
for i, (num, title, desc) in enumerate(problems):
    col, row = i % 2, i // 2
    x = sx2 + col * (cw + gx); y = sy2 + row * (ch + gy)
    card(s, x, y, cw, ch)
    icon_circle(s, x + Inches(0.2), y + Inches(0.2), Inches(0.5), num, sz=18)
    add_text(s, x + Inches(0.85), y + Inches(0.15), cw - Inches(1.0), Inches(0.5), title, sz=15, bold=True, color=GREEN_DARK)
    add_text(s, x + Inches(0.3), y + Inches(0.8), cw - Inches(0.6), Inches(1.5), desc, sz=12)
footer(s, 5)

# ═══ SLIDE 6: ALASAN PROYEK ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Alasan Proyek Perlu Dikembangkan")
reasons = [("Mendigitalisasi Proses Bank Sampah","Mengubah proses manual menjadi platform digital yang terstruktur dan efisien."),("Memberikan Edukasi Pemilahan Sampah","Memberikan informasi langsung kepada masyarakat tentang jenis-jenis sampah dan nilainya."),("Memberikan Nilai Ekonomis","Sistem penimbangan dan pembayaran yang transparan untuk sampah yang dikelola."),("Meningkatkan Partisipasi Masyarakat","Mendorong masyarakat aktif memilah sampah melalui sistem reward yang jelas."),("Tools bagi Pengelola Bank Sampah","Mengelola operasional secara efisien dan terdokumentasi dengan baik.")]
sx3, sy3 = Inches(0.8), Inches(1.3)
rw, rh = Inches(12), Inches(1.05)
rg = Inches(0.12)
for i, (title, desc) in enumerate(reasons):
    y = sy3 + i * (rh + rg)
    card(s, sx3, y, rw, rh)
    icon_circle(s, sx3 + Inches(0.15), y + Inches(0.275), Inches(0.5), str(i+1), sz=18)
    add_text(s, sx3 + Inches(0.85), y + Inches(0.1), Inches(4), Inches(0.4), title, sz=14, bold=True, color=GREEN_DARK)
    add_text(s, sx3 + Inches(5.2), y + Inches(0.15), Inches(6.5), Inches(0.8), desc, sz=12, color=GRAY_TEXT)
footer(s, 6)

# ═══ SLIDE 7: TUJUAN PENGEMBANGAN ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Tujuan Pengembangan Aplikasi")
goals = [
    "Menyediakan platform digital yang menghubungkan masyarakat dengan pengelola bank sampah",
    "Memberikan edukasi pemilahan sampah melalui informasi jenis sampah dan nilainya",
    "Memfasilitasi proses penjadwalan penjemputan sampah dari rumah masyarakat",
    "Menyediakan sistem penimbangan dan kalkulasi harga sampah yang transparan",
    "Mengintegrasikan dompet digital untuk pembayaran hasil setoran sampah",
    "Menyediakan dashboard riwayat transaksi untuk tracking dan transparansi",
    "Mendorong masyarakat aktif memilah dan menyetor sampah melalui sistem reward",
]
sy = Inches(1.3)
for i, goal in enumerate(goals):
    y = sy + i * Inches(0.78)
    icon_circle(s, Inches(1.0), y, Inches(0.5), str(i+1), sz=16)
    add_text(s, Inches(1.7), y + Inches(0.05), Inches(10.5), Inches(0.7), goal, sz=14)
    if i < len(goals) - 1:
        add_shape(s, MSO_SHAPE.RECTANGLE, Inches(1.25), y + Inches(0.55), Inches(0.01), Inches(0.25), GREEN_LIGHT)
footer(s, 7)

# ═══ SLIDE 8: TARGET PENGGUNA ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Target Pengguna & Manfaat")
# Card 1 - Masyarakat
card(s, Inches(0.8), Inches(1.3), Inches(5.8), Inches(5.5))
add_text(s, Inches(1.0), Inches(1.5), Inches(5.4), Inches(0.5), "Masyarakat (Penyetor Sampah)", sz=18, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
add_shape(s, MSO_SHAPE.RECTANGLE, Inches(2.3), Inches(2.05), Inches(3), Inches(0.03), ACCENT)
add_multi(s, Inches(1.2), Inches(2.3), Inches(5.2), Inches(4), [
    ("Profil:", True),
    ("- Masyarakat umum, usia 17+ tahun", False),
    ("- Melek teknologi (smartphone)", False),
    ("- Ingin penghasilan tambahan", False),
    ("", False),
    ("Manfaat:", True),
    ("- Edukasi pemilahan sampah", False),
    ("- Pendapatan dari sampah", False),
    ("- Penjemputan dari rumah", False),
    ("- Tracking riwayat setoran", False),
    ("- Dompet digital transparan", False),
], sz=13, sp=4)
# Card 2 - Petugas
card(s, Inches(6.9), Inches(1.3), Inches(5.8), Inches(5.5))
add_text(s, Inches(7.1), Inches(1.5), Inches(5.4), Inches(0.5), "Pengelola Bank Sampah (Petugas)", sz=18, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
add_shape(s, MSO_SHAPE.RECTANGLE, Inches(8.4), Inches(2.05), Inches(3), Inches(0.03), ACCENT)
add_multi(s, Inches(7.3), Inches(2.3), Inches(5.2), Inches(4), [
    ("Profil:", True),
    ("- Operator & petugas bank sampah", False),
    ("- Bertanggung jawab verifikasi", False),
    ("- Menggunakan aplikasi desktop", False),
    ("", False),
    ("Manfaat:", True),
    ("- Dashboard operasional", False),
    ("- Manajemen penjemputan", False),
    ("- Sistem penimbangan digital", False),
    ("- Laporan transaksi otomatis", False),
    ("- Manajemen harga sampah", False),
], sz=13, sp=4)
footer(s, 8)

# ═══ SLIDE 9: FITUR UTAMA ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Fitur Utama Aplikasi")
features = [
    ("F1","Registrasi & Manajemen Pengguna","Pendaftaran akun, login, profil, role-based access"),
    ("F2","Edukasi Pemilahan Sampah","Kategori sampah, harga per kg, panduan pemilahan"),
    ("F3","Penjadwalan Jemput Sampah","Pilih jenis sampah, jadwal, alamat, status"),
    ("F4","Sistem Penimbangan","Input berat aktual, kalkulasi harga otomatis"),
    ("F5","Dompet Digital & Pembayaran","Saldo, pembayaran otomatis, riwayat, penarikan"),
    ("F6","Dashboard Riwayat Transaksi","Riwayat setoran, pembayaran, statistik"),
    ("F7","Manajemen Operasional (Desktop)","Dashboard petugas, laporan, manajemen harga"),
]
sy = Inches(1.25)
fw = Inches(12); fh = Inches(0.78); fg = Inches(0.07)
for i, (code, title, desc) in enumerate(features):
    y = sy + i * (fh + fg)
    card(s, Inches(0.8), y, fw, fh)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y + Inches(0.12), Inches(0.8), Inches(0.5), GREEN_MAIN)
    tb = s.shapes.add_textbox(Inches(1.0), y + Inches(0.12), Inches(0.8), Inches(0.5))
    tf = tb.text_frame; p = tf.paragraphs[0]; p.text = code
    p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER; p.font.name = "Segoe UI"
    add_text(s, Inches(2.0), y + Inches(0.08), Inches(4), Inches(0.65), title, sz=14, bold=True, color=GREEN_DARK)
    add_text(s, Inches(6.2), y + Inches(0.08), Inches(6.3), Inches(0.65), desc, sz=12, color=GRAY_TEXT)
footer(s, 9)

# ═══ SLIDE 10: ALUR PENGGUNAAN ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Alur Penggunaan Aplikasi")
# Left - Masyarakat
card(s, Inches(0.5), Inches(1.3), Inches(6), Inches(5.5))
add_text(s, Inches(0.7), Inches(1.4), Inches(5.6), Inches(0.5), "Alur Masyarakat (Mobile App)", sz=16, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
add_shape(s, MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(1.95), Inches(4), Inches(0.03), ACCENT)
masyarati_flow = ["Registrasi Akun","Login","Lihat Edukasi Sampah","Buat Penjemputan","Pilih Jenis & Estimasi Berat","Pilih Jadwal & Alamat","Tunggu Penjemputan","Petugas Menimbang","Verifikasi Hasil","Pembayaran ke Dompet","Lihat Riwayat"]
fy = Inches(2.15)
for i, step in enumerate(masyarati_flow):
    y = fy + i * Inches(0.39)
    icon_circle(s, Inches(0.9), y, Inches(0.3), str(i+1), sz=10)
    add_text(s, Inches(1.35), y + Inches(0.0), Inches(4.8), Inches(0.35), step, sz=12)
    if i < len(masyarati_flow) - 1:
        add_shape(s, MSO_SHAPE.RECTANGLE, Inches(1.04), y + Inches(0.3), Inches(0.02), Inches(0.08), GREEN_LIGHT)
# Right - Petugas
card(s, Inches(6.8), Inches(1.3), Inches(6), Inches(5.5))
add_text(s, Inches(7.0), Inches(1.4), Inches(5.6), Inches(0.5), "Alur Petugas (Desktop App)", sz=16, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
add_shape(s, MSO_SHAPE.RECTANGLE, Inches(7.8), Inches(1.95), Inches(4), Inches(0.03), ACCENT)
petugas_flow = ["Login Petugas","Dashboard Operasional","Lihat Penjemputan Terjadwal","Pilih Penjemputan","Datang ke Lokasi","Timbang Sampah","Input Berat Aktual","Sistem Kalkulasi Harga","Konfirmasi Pembayaran","Transfer ke Dompet","Update Status = Selesai","Lihat Laporan"]
fy2 = Inches(2.15)
for i, step in enumerate(petugas_flow):
    y = fy2 + i * Inches(0.36)
    icon_circle(s, Inches(7.2), y, Inches(0.3), str(i+1), sz=10)
    add_text(s, Inches(7.65), y, Inches(4.8), Inches(0.35), step, sz=12)
    if i < len(petugas_flow) - 1:
        add_shape(s, MSO_SHAPE.RECTANGLE, Inches(7.34), y + Inches(0.3), Inches(0.02), Inches(0.05), GREEN_LIGHT)
footer(s, 10)

# ═══ SLIDE 11: RENCANA OOP - CLASS DIAGRAM ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Rencana OOP - Class & Attribute")
classes = [
    ("User (abstract)","id, nama, email, password, alamat, noTelp","getRole() [abstract]","GREEN_DARK"),
    ("Masyarakat extends User","saldoDompet, poinReward, riwayatSetoran","buatPenjemputan(), verifikasiTimbangan(), tarikSaldo()","GREEN_MAIN"),
    ("Petugas extends User","nip, shift, areaTugas","prosesPenjemputan(), inputTimbangan(), konfirmasiPembayaran()","GREEN_MAIN"),
    ("Sampah","id, kategori, nama, hargaPerKg, deskripsi","getHarga(), getKategori()","GREEN_LIGHT"),
    ("Penjemputan","id, masyarakat, tanggal, status, alamat","addDetail(), calculateTotal(), updateStatus()","GREEN_LIGHT"),
    ("DetailPenjemputan","penjemputan, sampah, estimasiBerat, beratAktual","calculateSubtotal()","GREEN_LIGHT"),
    ("Transaksi","id, penjemputan, totalHarga, tanggal, status","processPayment(), getDetail()","ACCENT"),
    ("DompetDigital","id, masyarakat, saldo, riwayatTransaksi","topUp(), withdraw(), getSaldo()","ACCENT"),
    ("Notifikasi","id, penerima, pesan, tanggal, statusBaca","send(), markAsRead()","ACCENT"),
]
cw = Inches(4.0); ch = Inches(1.55)
gx, gy = Inches(0.15), Inches(0.15)
sx, sy = Inches(0.5), Inches(1.15)
for i, (name, attrs, methods, color_name) in enumerate(classes):
    col, row = i % 3, i // 3
    x = sx + col * (cw + gx)
    y = sy + row * (ch + gy)
    c = GREEN_DARK if color_name == "GREEN_DARK" else (GREEN_MAIN if color_name == "GREEN_MAIN" else (GREEN_LIGHT if color_name == "GREEN_LIGHT" else ACCENT))
    card(s, x, y, cw, ch, GREEN_BG)
    add_shape(s, MSO_SHAPE.RECTANGLE, x, y, cw, Inches(0.35), c)
    add_text(s, x + Inches(0.1), y + Inches(0.02), cw - Inches(0.2), Inches(0.35), name, sz=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.1), y + Inches(0.42), cw - Inches(0.2), Inches(0.5), "+ " + attrs, sz=9, color=DARK_TEXT)
    add_text(s, x + Inches(0.1), y + Inches(1.0), cw - Inches(0.2), Inches(0.5), "+ " + methods, sz=9, color=GREEN_DARK)
footer(s, 11)

# ═══ SLIDE 12: RENCANA OOP - KONSEP ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Rencana OOP - Konsep yang Diterapkan")
concepts = [
    ("Encapsulation","Semua attribute private, akses via getter & setter dengan validasi","private double saldoDompet;\npublic double getSaldoDompet() { return saldoDompet; }\npublic void setSaldoDompet(double s) { this.saldoDompet = s; }"),
    ("Inheritance","Masyarakat & Petugas mewarisi dari User (abstract class)","User (abstract)\n  |-- Masyarakat\n  |-- Petugas"),
    ("Polymorphism","Method overriding (getRole) & overloading (buatPenjemputan)","@Override\npublic String getRole() { return \"MASYARAKAT\"; }\n\npublic Penjemputan buatPenjemputan(List<Sampah> s, LocalDate t)\npublic Penjemputan buatPenjemputan(List<Sampah> s, LocalDate t, String c)"),
    ("Constructor","Setiap class utama memiliki constructor untuk inisialisasi","public Masyarakat(String nama, String email,\n    String password, String alamat, String noTelp) { super(...); }"),
    ("Object & Method","Object dibuat dari class, method untuk behavior","Masyarakat m = new Masyarakat(...);\nPenjemputan p = m.buatPenjemputan(sampah, tanggal);"),
    ("Relasi Class","Association antar class (one-to-many, one-to-one)","Masyarakat 1--* Penjemputan\nPenjemputan 1--* DetailPenjemputan\nMasyarakat 1--1 DompetDigital"),
]
cw = Inches(6.0); ch = Inches(1.75)
gx, gy = Inches(0.15), Inches(0.15)
sx, sy = Inches(0.5), Inches(1.15)
for i, (title, desc, code) in enumerate(concepts):
    col, row = i % 2, i // 2
    x = sx + col * (cw + gx)
    y = sy + row * (ch + gy)
    card(s, x, y, cw, ch)
    add_text(s, x + Inches(0.2), y + Inches(0.05), cw - Inches(0.4), Inches(0.35), title, sz=14, bold=True, color=GREEN_DARK)
    add_text(s, x + Inches(0.2), y + Inches(0.4), cw - Inches(0.4), Inches(0.4), desc, sz=11, color=GRAY_TEXT)
    add_shape(s, MSO_SHAPE.RECTANGLE, x + Inches(0.2), y + Inches(0.85), cw - Inches(0.4), Inches(0.8), RGBColor(0xF5, 0xF5, 0xF5))
    add_text(s, x + Inches(0.3), y + Inches(0.88), cw - Inches(0.6), Inches(0.75), code, sz=9, color=DARK_TEXT)
footer(s, 12)

# ═══ SLIDE 13: TAHAPAN PENGEMBANGAN ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Rencana Tahapan Pengembangan hingga UAS")
stages = [
    ("Tahap 1","Persiapan & Desain (Pra-UTS)","PRD, wireframe Figma, rancangan class, setup NetBeans", True),
    ("Tahap 2","Prototype (UTS)","Class dasar, UI registrasi & login, demo 1 alur fungsi", True),
    ("Tahap 3","Pengembangan Inti (Pasca-UTS)","Database, penjadwalan, penimbangan, dompet digital, dashboard", False),
    ("Tahap 4","Aplikasi Desktop Petugas","Dashboard operasional, manajemen penjemputan, laporan", False),
    ("Tahap 5","Integrasi & Polish (UAS)","Integrasi mobile+desktop, fitur tambahan, testing, finalisasi", False),
]
sy = Inches(1.3)
for i, (stage, title, desc, done) in enumerate(stages):
    y = sy + i * Inches(1.12)
    card(s, Inches(0.8), y, Inches(12), Inches(1.0))
    # stage badge
    add_shape(s,
 MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y + Inches(0.15), Inches(1.2), Inches(0.7), GREEN_MAIN)
    tb = s.shapes.add_textbox(Inches(1.0), y + Inches(0.15), Inches(1.2), Inches(0.7))
    tf = tb.text_frame; p = tf.paragraphs[0]; p.text = stage
    p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER; p.font.name = "Segoe UI"
    # title + desc
    add_text(s, Inches(2.4), y + Inches(0.1), Inches(7), Inches(0.4), title, sz=14, bold=True, color=GREEN_DARK)
    add_text(s, Inches(2.4), y + Inches(0.5), Inches(7), Inches(0.4), desc, sz=11, color=GRAY_TEXT)
    # status
    status_text = "SELESAI" if done else "AKAN DATANG"
    status_color = ACCENT if done else ORANGE
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), y + Inches(0.25), Inches(2), Inches(0.5), status_color)
    tb2 = s.shapes.add_textbox(Inches(10.5), y + Inches(0.25), Inches(2), Inches(0.5))
    tf2 = tb2.text_frame; p2 = tf2.paragraphs[0]; p2.text = status_text
    p2.font.size = Pt(10); p2.font.bold = True; p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER; p2.font.name = "Segoe UI"
footer(s, 13)

# ═══ SLIDE 14: PEMBAGIAN TUGAS ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Pembagian Tugas Anggota Kelompok")
tasks = [
    ("Anggota 1","Arsitektur & Backend Core",["Rancangan class diagram","Implementasi class dasar (User, Masyarakat, Petugas)","Inheritance & polymorphism","Koneksi database"]),
    ("Anggota 2","Fitur Transaksi & Dompet",["Implementasi Penjemputan & DetailPenjemputan","Class Transaksi & DompetDigital","Sistem penimbangan & kalkulasi harga","Integrasi pembayaran"]),
    ("Anggota 3","UI/UX & Edukasi",["Desain & implementasi UI (mobile & desktop)","Fitur edukasi pemilahan sampah","Dashboard riwayat transaksi","Sistem notifikasi"]),
]
cw = Inches(4.0); ch = Inches(5.2); gap = Inches(0.25)
tw = cw * 3 + gap * 2; sx = (SLIDE_W - tw) // 2; cy = Inches(1.3)
for i, (member, role, items) in enumerate(tasks):
    x = sx + i * (cw + gap)
    card(s, x, cy, cw, ch)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, cy, cw, Inches(0.8), GREEN_MAIN)
    add_text(s, x + Inches(0.1), cy + Inches(0.1), cw - Inches(0.2), Inches(0.35), member, sz=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.1), cy + Inches(0.42), cw - Inches(0.2), Inches(0.35), role, sz=11, color=GREEN_LIGHT, align=PP_ALIGN.CENTER)
    for j, item in enumerate(items):
        y = cy + Inches(1.05) + j * Inches(0.85)
        icon_circle(s, x + Inches(0.2), y, Inches(0.3), str(j+1), sz=10)
        add_text(s, x + Inches(0.65), y + Inches(0.0), cw - Inches(0.8), Inches(0.8), item, sz=11)
footer(s, 14)

# ═══ SLIDE 15: TARGET & FITUR TAMBAHAN UAS ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, WHITE); header(s, "Target Pengembangan & Fitur Tambahan UAS")
add_text(s, Inches(0.8), Inches(1.2), Inches(6), Inches(0.4), "Fitur Tambahan yang Direncanakan:", sz=16, bold=True, color=GREEN_DARK)
extras = [
    "Sistem Poin Reward - poin setiap setoran, tukar voucher",
    "Leaderboard Kontribusi - ranking masyarakat berdasarkan setoran",
    "Edukasi Interaktif - kuis pemilahan sampah dengan reward",
    "Notifikasi Push - pengingat jadwal & update status",
    "Laporan Lingkungan - statistik dampak (kg sampah, CO2)",
    "Multi-Cabang Bank Sampah - dukungan multiple lokasi",
    "Export Laporan (PDF/Excel) - transaksi masyarakat & petugas",
]
sy = Inches(1.65)
for i, item in enumerate(extras):
    y = sy + i * Inches(0.55)
    icon_circle(s, Inches(1.0), y, Inches(0.35), "+" if False else str(i+1), sz=12, bg=ACCENT)
    add_text(s, Inches(1.5), y + Inches(0.02), Inches(5.5), Inches(0.5), item, sz=13)
# Right side - target UAS
card(s, Inches(7.5), Inches(1.2), Inches(5.3), Inches(5.5))
add_text(s, Inches(7.7), Inches(1.35), Inches(4.9), Inches(0.5), "Hasil Akhir UAS:", sz=16, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
add_shape(s, MSO_SHAPE.RECTANGLE, Inches(8.5), Inches(1.9), Inches(3.3), Inches(0.03), ACCENT)
add_multi(s, Inches(7.8), Inches(2.1), Inches(4.8), Inches(4), [
    ("- Aplikasi mobile & desktop fully functional", False),
    ("- Database terintegrasi", False),
    ("- Minimal 5 alur fungsi end-to-end", False),
    ("- Dokumentasi lengkap (class diagram, ERD)", False),
    ("- User manual", False),
    ("- Presentasi final & demo", False),
    ("", False),
    ("Timeline:", True),
    ("UTS: Prototype & Rencana", True),
    ("Pasca-UTS: Pengembangan inti", False),
    ("Menuju UAS: Integrasi & polish", False),
    ("UAS: Presentasi final", True),
], sz=12, sp=6)
footer(s, 15)

# ═══ SLIDE 16: CLOSING / TERIMA KASIH ═══
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s, GREEN_DARK)
add_shape(s, MSO_SHAPE.OVAL, Inches(-3), Inches(-3), Inches(7), Inches(7), GREEN_MAIN)
add_shape(s, MSO_SHAPE.OVAL, Inches(10), Inches(4), Inches(8), Inches(8), GREEN_MAIN)
add_shape(s, MSO_SHAPE.OVAL, Inches(8), Inches(-2), Inches(5), Inches(5), GREEN_LIGHT)
logo_resik(s, Inches(5.4), Inches(1.2), Inches(2.5), Inches(0.9))
logo_r(s, Inches(5.9), Inches(2.5), Inches(1.5), Inches(1.5))
add_text(s, Inches(1), Inches(4.2), Inches(11.333), Inches(0.8), "Terima Kasih", sz=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_shape(s, MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(5.0), Inches(2.3), Inches(0.04), ACCENT)
add_text(s, Inches(1), Inches(5.2), Inches(11.333), Inches(0.5), "Pertanyaan & Diskusi", sz=20, color=GREEN_LIGHT, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(5.9), Inches(11.333), Inches(0.5), "RESIK - Bank Sampah Digital", sz=16, color=GRAY_TEXT, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(6.4), Inches(11.333), Inches(0.5), "Pemrograman Berorientasi Objek (PBO) - UTS 2026", sz=12, color=GRAY_TEXT, align=PP_ALIGN.CENTER)

# ═══ SAVE ═══
output_path = os.path.join(BASE_DIR, "..", "RESIK-Presentasi-UTS.pptx")
prs.save(output_path)
print(f"PPT saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")
